"""测试 XLSX 导出。"""
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from fapiao_helper.export import xlsx


def _sample_invoices() -> list[dict]:
    return [
        {"id": 1, "invoice_no": "INV001", "invoice_date": "2024-03-15",
         "buyer": "买方A", "seller": "卖方A", "amount": 1000.00, "tax": 60.00, "total": 1060.00,
         "category": "餐饮住宿", "project": "项目A", "status": "done", "needs_review": 0},
        {"id": 2, "invoice_no": "INV002", "invoice_date": "2024-03-16",
         "buyer": "买方B", "seller": "卖方B", "amount": 2000.00, "tax": 120.00, "total": 2120.00,
         "category": "交通", "project": "项目B", "status": "pending", "needs_review": 1},
    ]


def test_export_creates_file(tmp_path: Path):
    out = xlsx.export_xlsx(_sample_invoices(), tmp_path)
    assert out.exists()


def test_export_has_4_sheets(tmp_path: Path):
    out = xlsx.export_xlsx(_sample_invoices(), tmp_path)
    wb = load_workbook(out)
    assert len(wb.sheetnames) == 4


def test_export_empty_invoices(tmp_path: Path):
    out = xlsx.export_xlsx([], tmp_path)
    assert out.exists()
    wb = load_workbook(out)
    assert len(wb.sheetnames) == 4


def test_export_data_rows(tmp_path: Path):
    out = xlsx.export_xlsx(_sample_invoices(), tmp_path)
    wb = load_workbook(out)
    detail_ws = None
    for name in wb.sheetnames:
        if "明细" in name:
            detail_ws = wb[name]
            break
    assert detail_ws is not None
    assert detail_ws.max_row >= 4  # 表头 + 2 数据 + 合计


def test_export_custom_filename(tmp_path: Path):
    out = xlsx.export_xlsx(_sample_invoices(), tmp_path, filename="custom.xlsx")
    assert out.name == "custom.xlsx"
