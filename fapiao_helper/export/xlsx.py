"""XLSX 导出:4 个 Sheet(汇总/按项目/按类型/明细)。"""
from __future__ import annotations

from pathlib import Path
from typing import Any

from ..paths import ensure_dir


def _to_float(v: Any) -> float:
    try:
        return round(float(v or 0), 2)
    except (TypeError, ValueError):
        return 0.0


def export_xlsx(invoices: list[dict], export_dir: Path | str, filename: str = "invoices.xlsx") -> Path:
    """生成 4 Sheet 的 XLSX,返回文件路径。"""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    export_dir = Path(export_dir)
    ensure_dir(export_dir)
    out_path = export_dir / filename

    wb = Workbook()
    bold = Font(bold=True)

    # ---- Sheet 1: 汇总(按 status+project 分组) ----
    ws1 = wb.active
    ws1.title = "汇总"
    headers1 = ["状态", "项目", "类型", "数量", "金额合计", "税额合计", "价税合计"]
    ws1.append(headers1)
    for c in ws1[1]:
        c.font = bold

    groups: dict[tuple, dict] = {}
    for inv in invoices:
        key = (inv.get("status", ""), inv.get("project", ""))
        g = groups.setdefault(key, {"count": 0, "amount": 0.0, "tax": 0.0, "total": 0.0, "cats": set()})
        g["count"] += 1
        g["amount"] += _to_float(inv.get("amount"))
        g["tax"] += _to_float(inv.get("tax"))
        g["total"] += _to_float(inv.get("total"))
        cat = inv.get("category") or "未分类"
        g["cats"].add(cat)

    if groups:
        for (status, project), g in sorted(groups.items()):
            ws1.append([
                status, project, "/".join(sorted(g["cats"])),
                g["count"], round(g["amount"], 2), round(g["tax"], 2), round(g["total"], 2),
            ])
        # 合计行
        total_count = sum(g["count"] for g in groups.values())
        total_amount = sum(g["amount"] for g in groups.values())
        total_tax = sum(g["tax"] for g in groups.values())
        total_total = sum(g["total"] for g in groups.values())
        ws1.append(["合计", "", "", total_count, round(total_amount, 2), round(total_tax, 2), round(total_total, 2)])
        for c in ws1[ws1.max_row]:
            c.font = bold
    else:
        ws1.append(["无数据", "", "", 0, 0, 0, 0])
    _set_widths(ws1)

    # ---- Sheet 2: 按项目 ----
    ws2 = wb.create_sheet("按项目")
    _write_grouped(ws2, invoices, "project", bold)

    # ---- Sheet 3: 按类型 ----
    ws3 = wb.create_sheet("按类型")
    _write_grouped(ws3, invoices, "category", bold)

    # ---- Sheet 4: 明细 ----
    ws4 = wb.create_sheet("明细")
    headers4 = ["ID", "发票号", "开票日期", "购买方", "销售方", "金额", "税额", "价税合计", "类型", "项目", "状态", "需复核"]
    ws4.append(headers4)
    for c in ws4[1]:
        c.font = bold
    if invoices:
        sum_amount = sum_tax = sum_total = 0.0
        for inv in invoices:
            amt = _to_float(inv.get("amount"))
            tax = _to_float(inv.get("tax"))
            tot = _to_float(inv.get("total"))
            sum_amount += amt
            sum_tax += tax
            sum_total += tot
            ws4.append([
                inv.get("id", ""), inv.get("invoice_no", ""), inv.get("invoice_date", ""),
                inv.get("buyer", ""), inv.get("seller", ""),
                amt, tax, tot,
                inv.get("category", ""), inv.get("project", ""),
                inv.get("status", ""), "是" if inv.get("needs_review") else "否",
            ])
        ws4.append([f"合计(共{len(invoices)}条)", "", "", "", "", round(sum_amount, 2), round(sum_tax, 2), round(sum_total, 2), "", "", "", ""])
        for c in ws4[ws4.max_row]:
            c.font = bold
    else:
        ws4.append(["无数据"] * len(headers4))
    _set_widths(ws4)

    wb.save(str(out_path))
    return out_path


def _write_grouped(ws, invoices: list[dict], field: str, bold) -> None:
    """单字段分组的通用写入。"""
    headers = ["项目" if field == "project" else "类型", "数量", "金额", "税额", "价税合计"]
    ws.append(headers)
    for c in ws[1]:
        c.font = bold

    groups: dict[str, dict] = {}
    for inv in invoices:
        key = inv.get(field) or "未分类"
        g = groups.setdefault(key, {"count": 0, "amount": 0.0, "tax": 0.0, "total": 0.0})
        g["count"] += 1
        g["amount"] += _to_float(inv.get("amount"))
        g["tax"] += _to_float(inv.get("tax"))
        g["total"] += _to_float(inv.get("total"))

    if groups:
        for key, g in sorted(groups.items()):
            ws.append([key, g["count"], round(g["amount"], 2), round(g["tax"], 2), round(g["total"], 2)])
        total_count = sum(g["count"] for g in groups.values())
        total_amount = sum(g["amount"] for g in groups.values())
        total_tax = sum(g["tax"] for g in groups.values())
        total_total = sum(g["total"] for g in groups.values())
        ws.append(["合计", total_count, round(total_amount, 2), round(total_tax, 2), round(total_total, 2)])
        for c in ws[ws.max_row]:
            c.font = bold
    else:
        ws.append(["无数据", 0, 0, 0, 0])
    _set_widths(ws)


def _set_widths(ws) -> None:
    """按列内容最大长度设置列宽。"""
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            try:
                ln = len(str(cell.value or ""))
                if ln > max_len:
                    max_len = ln
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = min(max_len + 2, 40)
